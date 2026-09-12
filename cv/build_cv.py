"""Build a one-page editable CV template for the MyQuant portfolio."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "CV_Template.docx"

NAVY = RGBColor(25, 61, 98)
GREY = RGBColor(80, 80, 80)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def text(p, value, size=9.2, bold=False, color=None, italic=False):
    run = p.add_run(value)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


def heading(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    text(p, title.upper(), size=10, bold=True, color=NAVY)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "193D62")
    borders.append(bottom)
    p_pr.append(borders)


def role_line(doc, left, right):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(5.25)
    table.columns[1].width = Inches(1.65)
    for i, value in enumerate((left, right)):
        cell = table.cell(0, i)
        set_cell_margins(cell, 0, 0, 0, 0)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.RIGHT
        text(p, value, size=8.8, bold=(i == 0))


def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.17)
    p.paragraph_format.first_line_indent = Inches(-0.10)
    p.paragraph_format.space_after = Pt(0.5)
    p.paragraph_format.line_spacing = 1.0
    text(p, value, size=8.65)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.38)
    sec.bottom_margin = Inches(0.36)
    sec.left_margin = Inches(0.58)
    sec.right_margin = Inches(0.58)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    styles["Normal"].paragraph_format.space_after = Pt(0)
    styles["List Bullet"].font.name = "Aptos"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    text(p, "[YOUR NAME]", size=17, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    text(p, "Oxford, UK  |  [phone]  |  [email]  |  [LinkedIn]  |  [GitHub]", size=8.5, color=GREY)

    heading(doc, "Education")
    role_line(doc, "University of Oxford — MEng Engineering Science", "Expected 2030")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    text(p, "Incoming first-year student. Relevant areas: probability, computing, electronics and applied mathematics. ", size=8.65)
    text(p, "Replace this line with verified grades, awards or modules only.", size=8.65, italic=True, color=GREY)

    heading(doc, "Selected projects")
    role_line(doc, "ETF Momentum Research | Python, pandas, NumPy", "2011–2026")
    bullet(doc, "Built a reproducible 24-ETF, monthly 12–1 momentum backtest with a one-trading-day signal lag, explicit 10 bps one-way trading costs, research/out-of-sample separation and cached raw data.")
    bullet(doc, "Benchmarked equal-weight, momentum, inverse-volatility and 200-day SPY regime-filter variants; reported CAGR, volatility, Sharpe, drawdown and turnover, while documenting selection, execution and capacity limitations.")
    role_line(doc, "Option Pricing & Hedging Lab | Python, Monte Carlo", "2026")
    bullet(doc, "Implemented Black–Scholes prices and Greeks, Monte Carlo valuation and finite-difference checks; unit-tested put–call parity, convergence and numerical Greek error bounds.")
    bullet(doc, "Simulated daily and weekly discrete delta hedging under volatility misspecification and jump risk, focusing on P&L distribution and model-risk interpretation rather than a trading claim.")
    role_line(doc, "Applied Materials Equity Research | Fundamental analysis", "2026")
    bullet(doc, "Prepared a watchlist stock-pitch memo using FY2025 10-K and Q3 FY2026 disclosures; linked semiconductor-equipment process complexity, advanced packaging and service revenue to catalysts and disconfirming risks.")

    heading(doc, "Experience & leadership")
    role_line(doc, "[Organisation / team] — [Role]", "[Dates]")
    bullet(doc, "[Use one evidence-based bullet: the task, your personal contribution, the method used and a verified outcome.]")
    role_line(doc, "[Organisation / team] — [Role]", "[Dates]")
    bullet(doc, "[Add a second genuine experience, engineering build, competition or leadership example. Do not invent metrics.]")

    heading(doc, "Skills & interests")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    text(p, "Technical: ", size=8.65, bold=True)
    text(p, "Python, NumPy, pandas, Git; [add SQL only if you can use it live].  ", size=8.65)
    text(p, "Finance: ", size=8.65, bold=True)
    text(p, "quantitative research, derivatives simulation, company research.  ", size=8.65)
    text(p, "Interests: ", size=8.65, bold=True)
    text(p, "[one genuine interest that can start a conversation].", size=8.65)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "Template note: personalise every bracketed field and be ready to explain every completed bullet for two minutes.", size=7.6, italic=True, color=GREY)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
